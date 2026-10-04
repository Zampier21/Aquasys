"""
Gera o resumo expandido do AquaSys para o SAPIEN 2026 a partir do modelo
oficial do evento.

    python gerar_resumo_sapien.py

O documento nao e' montado do zero: cada paragrafo novo e' uma copia de um
paragrafo do proprio modelo (titulo, autor, resumo, titulo de secao,
corpo, referencia), com o texto trocado. Assim fonte, recuo, entrelinha e
alinhamento sao os que a organizacao definiu, e nao uma imitacao deles.

As referencias sao lidas do Markdown do TCC, e nao redigitadas: o resumo
cita as mesmas obras, e uma divergencia entre os dois textos seria o tipo
de coisa que um avaliador nota.

Regras do modelo seguidas aqui:
  . resumo entre 50 e 100 palavras, paragrafo unico, sem citacao;
  . 3 a 5 palavras-chave separadas e finalizadas por ponto;
  . citacao no formato do modelo, (Autor, ano), no fim da frase;
  . titulo da obra em negrito nas referencias, como no exemplo do modelo.
"""

import copy
import os
import re
import unicodedata

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from docx.shared import Cm, Pt
from docx.text.paragraph import Paragraph

AQUI = os.path.dirname(os.path.abspath(__file__))
MODELO = r"C:\Users\wesll\Downloads\Sapien-2026-PESQUISA.docx"
TCC = os.path.join(AQUI, "..", "AquaSys_TCC.md")
FIGURA = os.path.join(AQUI, "..", "figura_arquitetura.png")
SAIDA = os.path.join(AQUI, "SAPIEN_2026_AquaSys_Resumo_Expandido.docx")

# ═══════════════════════════════════════════════════════
# CONTEÚDO
# ═══════════════════════════════════════════════════════
TITULO = ("AQUASYS: APLICATIVO MÓVEL PARA GESTÃO DE AQUÁRIOS E APOIO À "
          "MANUTENÇÃO PROFISSIONAL EM MODELO SAAS")

AUTORES = ["ZAMPIER, Weslley Cesar", "LEVANDOSKI, Brenda Lopes"]

RESUMO = (
    "O aquarismo ornamental exige controlar a qualidade da água e a "
    "compatibilidade entre espécies, conhecimentos que o iniciante "
    "raramente domina. Este trabalho apresenta o AquaSys, "
    "aplicativo móvel em modelo SaaS que interpreta os parâmetros da água "
    "conforme as espécies do aquário e orienta a composição do povoamento. "
    "A pesquisa seguiu o *Design Thinking* e um mapeamento sistemático da "
    "literatura. O sistema adota arquitetura cliente-servidor em três "
    "camadas, com as regras de negócio no servidor. O protótipo está "
    "funcional e verificado por 293 testes automatizados, e a usabilidade "
    "será validada por Avaliação Heurística."
)

PALAVRAS_CHAVE = ("Aquarismo. Aplicativo móvel. Arquitetura de software. "
                  "Qualidade da água. Software como serviço.")

# Cada item: ("h1" | "h2" | "p" | "figura", texto)
CORPO = [
    ("h1", "1 INTRODUÇÃO"),
    ("p", "O aquarismo ornamental é uma das atividades de lazer mais "
          "difundidas no mundo, com um comércio global que movimenta "
          "bilhões de dólares por ano e abrange milhares de espécies "
          "(Biondo; Burki, 2021). No Brasil, a aquicultura ornamental cresce "
          "sustentada pela diversidade da fauna nativa e pela demanda por "
          "espécies tropicais (Mendonça; Thomé, 2020)."),
    ("p", "Esse crescimento convive com a perda frequente de animais logo "
          "nas primeiras semanas, uma vez que manter um aquário equilibrado "
          "significa manter um ecossistema cujo funcionamento depende de "
          "variáveis químicas invisíveis a olho nu (Arana, 2004). Desvios "
          "discretos em pH, temperatura, amônia, nitrito e nitrato bastam "
          "para causar estresse, doença e morte, complexidade que o "
          "aquarista iniciante costuma descobrir apenas depois das primeiras "
          "perdas (Lins, 2021). Além disso, cada espécie tolera uma faixa "
          "própria desses parâmetros, definida pelo ambiente em que evoluiu "
          "(Baldisserotto, 2025). A convivência entre espécies, por sua vez, "
          "exige considerar em conjunto o porte, o comportamento e a faixa de "
          "parâmetros tolerada por cada uma (Mohammadi; Jafari, 2014)."),
    ("p", "As soluções existentes atacam apenas parte do problema. A "
          "automação por sensores é eficaz na coleta do dado, mas não "
          "orienta sobre o significado da medição nem trata da convivência "
          "entre espécies (Unoki; Candido, 2019; Mazziero *et al.*, 2023). "
          "Há ainda um público pouco atendido por essas soluções: o "
          "profissional que presta manutenção de aquários a domicílio. Em "
          "levantamento realizado na loja Aqualife Ecossistemas, em "
          "Guarapuava, observou-se que esse profissional registra cada visita "
          "em fichas de papel, que depois não podem ser consultadas nem "
          "comparadas com o atendimento anterior."),
    ("p", "Diante desse cenário, este trabalho tem como objetivo desenvolver "
          "e avaliar o AquaSys, aplicativo móvel distribuído no modelo "
          "*Software as a Service* (SaaS) que interpreta os parâmetros da "
          "água à luz das espécies presentes, orienta a composição do "
          "povoamento e substitui a ficha de papel por um histórico "
          "consultável. A seção seguinte descreve os métodos empregados e "
          "os resultados obtidos até o momento."),

    ("h1", "2 DESENVOLVIMENTO"),
    ("h2", "2.1 Materiais e métodos"),
    ("p", "A pesquisa é de natureza aplicada, por destinar-se à construção "
          "de um artefato para resolver um problema concreto, e de abordagem "
          "qualitativa, por buscar compreender as necessidades de um público "
          "específico e traduzi-las em requisitos (Prodanov; Freitas, 2013). "
          "A condução seguiu o *Design Thinking*, organizado em cinco "
          "etapas: empatia, definição, ideação, prototipação e teste "
          "(Lopes, 2023)."),
    ("p", "Na etapa de empatia, realizou-se observação direta em ambiente "
          "comercial de aquarismo e conversas com aquaristas de diferentes "
          "níveis de experiência. Paralelamente, conduziu-se um mapeamento "
          "sistemático da literatura, método que identifica, classifica e "
          "resume a produção científica sobre um tema (Kitchenham; Charters, "
          "2007). As buscas na IEEE *Xplore* e no Google Acadêmico, entre "
          "2015 e 2026, retornaram 880 registros, e oito estudos foram "
          "selecionados após a remoção de duplicados, a triagem por título "
          "e resumo e a leitura integral."),
    ("p", "O sistema adota arquitetura cliente-servidor em três camadas, "
          "apresentada na Figura 1. A opção por um monolito modular, em vez "
          "de microsserviços, segue a recomendação de começar por um "
          "monolito e extrair serviços apenas quando as fronteiras "
          "estiverem comprovadas pelo uso (Fowler, 2015). A comunicação "
          "segue o estilo REST, com requisições autocontidas e sem estado "
          "de sessão no servidor (Fielding, 2000). O isolamento entre as "
          "lojas é feito por multilocação por linha, em que todas "
          "compartilham as mesmas tabelas e cada registro identifica o seu "
          "proprietário (Chong; Carraro; Wolter, 2006)."),
    ("figura", "Figura 1 – Arquitetura em três camadas do AquaSys"),
    ("p", "A camada de apresentação é um aplicativo em Flutter, arcabouço "
          "que gera versões para Android e iOS a partir de uma base de "
          "código única (Flutter, 2025). A lógica de negócio é uma API em "
          "FastAPI, que valida a entrada e a saída por modelos declarativos "
          "antes de qualquer regra ser executada (FastAPI, 2024). Os dados "
          "residem em PostgreSQL, com integridade referencial e restrições "
          "de verificação declaradas no próprio esquema (PostgreSQL Global "
          "Development Group, 2025)."),
    ("p", "A verificação é feita por testes automatizados de unidade, de "
          "integração e de interface. A validação de usabilidade será "
          "conduzida por Avaliação Heurística, em que a interface é "
          "confrontada com dez princípios gerais de usabilidade e cada "
          "violação recebe um grau de severidade (Nielsen, 1994)."),

    ("h2", "2.2 Resultados parciais"),
    ("p", "O protótipo está funcional e cobre o ciclo de uso dos dois "
          "perfis, a loja assinante e o aquarista cliente. O servidor expõe "
          "48 rotas sobre 18 tabelas, e a regra de negócio concentra-se em "
          "serviços de domínio independentes do protocolo HTTP, o que "
          "permite testá-los sem servidor ativo."),
    ("p", "O principal resultado é o motor de avaliação de parâmetros. Em "
          "sua primeira versão, a faixa ideal era derivada do tipo declarado "
          "do aquário e produzia alarme falso sempre que o povoamento fugia "
          "da média do tipo. Na versão atual, a faixa é derivada dos peixes "
          "efetivamente presentes: como a mesma água serve a todos, o "
          "intervalo válido é a interseção das tolerâncias de cada espécie, "
          "ou seja, o maior dos mínimos e o menor dos máximos. Quando as "
          "exigências não se cruzam, o sistema não acusa a água, e sim "
          "aponta que o problema está na combinação de espécies. As "
          "mensagens identificam a espécie afetada, como em \"Esse parâmetro "
          "está errado pois o Neon Tetra vive em água mais ácida (pH até "
          "7,0), e a sua está em 7,3 (alcalina)\"."),
    ("p", "O motor de compatibilidade classifica cada combinação em quatro "
          "níveis e a converte em uma de três decisões: liberado, requer "
          "confirmação ou bloqueado. Combinações de risco previsível e "
          "irreversível, como a predação decorrente da diferença de porte, "
          "são bloqueadas com mensagem específica, do tipo \"Acará Bandeira "
          "adulto (15 cm) predará Neon Tetra (4 cm)\"."),
    ("p", "O próprio desenvolvimento ofereceu evidência para a decisão de "
          "concentrar as regras no servidor. Em uma versão intermediária, as "
          "faixas ideais existiam duplicadas no aplicativo e na API; as duas "
          "cópias divergiram, e o mesmo aquário passou a ser exibido como "
          "saudável em uma tela e problemático em outra. A correção consistiu "
          "em eliminar a cópia do aplicativo, episódio que ilustra que os "
          "atributos de qualidade decorrem da estrutura do sistema, e não do "
          "cuidado com cada trecho de código isolado (Bass; Clements; "
          "Kazman, 2021)."),
    ("p", "A separação das fotografias do catálogo em tabela própria manteve "
          "a listagem de espécies em 8,5 KB, sem nenhum byte de imagem, com "
          "miniaturas de 6,8 KB e revalidação por *ETag* que evita "
          "retransmitir imagem já obtida. A suíte de verificação reúne 293 "
          "testes automatizados, 270 na API e 23 no aplicativo, dos quais "
          "13 confirmam que uma loja não alcança dados de outra ao trocar "
          "identificadores na requisição."),

    ("h1", "3 CONSIDERAÇÕES FINAIS"),
    ("p", "O AquaSys demonstra a viabilidade técnica de um aplicativo que "
          "interpreta os parâmetros da água a partir das espécies presentes, "
          "e não de faixas fixas por tipo de aquário, distinção que o "
          "diferencia das soluções levantadas no mapeamento. Manter as regras "
          "de negócio no servidor trouxe dois ganhos concretos. Todas as telas "
          "passaram a exibir a mesma avaliação, por existir uma única "
          "implementação dela, e ampliar o catálogo ou corrigir uma faixa de "
          "parâmetro deixou de exigir atualização do aplicativo."),
    ("p", "Como próximas etapas, estão previstas a Avaliação Heurística do "
          "protótipo, a conclusão dos requisitos pendentes, entre eles o "
          "lembrete de troca parcial de água e a recuperação de senha por "
          "correio eletrônico, e a implantação do sistema em ambiente de "
          "produção."),
]

# Obras citadas, pelo início exato da entrada no TCC.
REFERENCIAS = [
    "ARANA,", "BALDISSEROTTO,", "BASS,", "BIONDO,", "CHONG,", "FASTAPI.",
    "FIELDING,", "FLUTTER.", "FOWLER,", "KITCHENHAM,", "LINS,", "LOPES,",
    "MAZZIERO,", "MENDONÇA,", "MOHAMMADI,", "NIELSEN, Jakob. ",
    "POSTGRESQL", "PRODANOV,", "UNOKI,",
]


# ═══════════════════════════════════════════════════════
# MONTAGEM
# ═══════════════════════════════════════════════════════
_MARCAS = re.compile(r"(\*\*.+?\*\*|\*[^*]+?\*)")


def ler_referencias():
    linhas = open(TCC, encoding="utf-8").read().split("## Referências")[1]
    linhas = [l.strip() for l in linhas.split("\n")]
    achadas = []
    for inicio in REFERENCIAS:
        casa = [l for l in linhas if l.startswith(inicio)]
        if len(casa) != 1:
            raise SystemExit("referência '%s': %d entradas" % (inicio, len(casa)))
        achadas.append(casa[0])
    return achadas


def limpar_runs(p_elm):
    for filho in list(p_elm):
        if filho.tag != qn("w:pPr"):
            p_elm.remove(filho)


def rpr_de(prototipo, indice=0):
    """Propriedades de fonte do n-ésimo run do parágrafo modelo."""
    runs = prototipo.findall(qn("w:r"))
    rpr = runs[indice].find(qn("w:rPr")) if runs else None
    return copy.deepcopy(rpr) if rpr is not None else None


def novo_run(p_elm, texto, rpr, negrito=None, italico=False, tamanho=None):
    r = p_elm.makeelement(qn("w:r"), {})
    rpr = copy.deepcopy(rpr) if rpr is not None else r.makeelement(qn("w:rPr"), {})
    for tag in ("w:b", "w:bCs", "w:i", "w:iCs"):
        for velho in rpr.findall(qn(tag)):
            rpr.remove(velho)
    # A ordem dos filhos do rPr é imposta pelo esquema: negrito e itálico
    # vêm logo depois da fonte, antes de tamanho e cor.
    pos = 1 if rpr.find(qn("w:rFonts")) is not None else 0
    if italico:
        rpr.insert(pos, rpr.makeelement(qn("w:i"), {}))
    if negrito:
        rpr.insert(pos, rpr.makeelement(qn("w:b"), {}))
    if tamanho:
        for velho in rpr.findall(qn("w:sz")) + rpr.findall(qn("w:szCs")):
            rpr.remove(velho)
        sz = rpr.makeelement(qn("w:sz"), {qn("w:val"): str(tamanho * 2)})
        rpr.append(sz)
    r.append(rpr)
    t = r.makeelement(qn("w:t"), {})
    t.text = texto
    t.set("{http://www.w3.org/XML/1998/namespace}space", "preserve")
    r.append(t)
    p_elm.append(r)


def escrever(p_elm, texto, rpr, negrito_base=False, tamanho=None,
             negrito_marcado=True):
    for pedaco in _MARCAS.split(texto):
        if not pedaco:
            continue
        if pedaco.startswith("**"):
            novo_run(p_elm, pedaco[2:-2], rpr, negrito=negrito_marcado,
                     tamanho=tamanho)
        elif pedaco.startswith("*"):
            novo_run(p_elm, pedaco[1:-1], rpr, negrito=negrito_base,
                     italico=True, tamanho=tamanho)
        else:
            novo_run(p_elm, pedaco, rpr, negrito=negrito_base, tamanho=tamanho)


def gerar():
    doc = Document(MODELO)
    ps = doc.paragraphs
    modelo = {
        "titulo": ps[0]._p, "autor": ps[1]._p, "vazio": ps[4]._p,
        "resumo": ps[5]._p, "chave": ps[7]._p, "h1": ps[9]._p,
        "corpo": ps[11]._p, "ref_titulo": ps[32]._p, "ref": ps[36]._p,
        "ref_vazio": ps[37]._p,
    }
    originais = [p._p for p in ps]
    corpo_xml = doc.element.body
    fim = corpo_xml.find(qn("w:sectPr"))
    novos = []

    def clonar(tipo):
        elm = copy.deepcopy(modelo[tipo])
        limpar_runs(elm)
        corpo_xml.insert(list(corpo_xml).index(fim), elm)
        novos.append(elm)
        return elm

    def vazio():
        clonar("vazio")

    # ─ título e autores ─
    escrever(clonar("titulo"), TITULO, rpr_de(modelo["titulo"]),
             negrito_base=True)
    for autor in AUTORES:
        escrever(clonar("autor"), autor, rpr_de(modelo["autor"]))
    vazio()

    # ─ resumo e palavras-chave: rótulo em negrito, texto em redondo ─
    p = clonar("resumo")
    novo_run(p, "RESUMO: ", rpr_de(modelo["resumo"], 0), negrito=True)
    escrever(p, RESUMO, rpr_de(modelo["resumo"], 1))
    vazio()
    p = clonar("chave")
    novo_run(p, "Palavras-chave: ", rpr_de(modelo["chave"], 0), negrito=True)
    escrever(p, PALAVRAS_CHAVE, rpr_de(modelo["chave"], 1))
    vazio()

    # ─ corpo ─
    rpr_corpo = rpr_de(modelo["corpo"])
    rpr_h1 = rpr_de(modelo["h1"])
    figuras = 0
    anterior = None
    for tipo, texto in CORPO:
        if tipo in ("h1", "h2"):
            if anterior is not None:
                vazio()
            escrever(clonar("h1"), texto, rpr_h1, negrito_base=True)
            vazio()
        elif tipo == "p":
            escrever(clonar("corpo"), texto, rpr_corpo)
        elif tipo == "figura":
            # Legenda acima e fonte abaixo, centralizadas, em corpo menor e
            # entrelinha simples, como pede a NBR 14724 para ilustrações.
            for papel, conteudo in (("legenda", texto), ("imagem", None),
                                    ("fonte", "Fonte: Os autores (2026).")):
                elm = clonar("corpo")
                par = Paragraph(elm, doc._body)
                pf = par.paragraph_format
                pf.first_line_indent = Cm(0)
                pf.line_spacing = 1.0
                pf.space_before = Pt(6 if papel == "legenda" else 0)
                pf.space_after = Pt(6 if papel == "fonte" else 0)
                par.alignment = WD_ALIGN_PARAGRAPH.CENTER
                if papel == "imagem":
                    par.add_run().add_picture(FIGURA, width=Cm(15.5))
                else:
                    escrever(elm, conteudo, rpr_corpo, tamanho=10)
            figuras += 1
        anterior = tipo

    # ─ referências ─
    vazio()
    p = clonar("ref_titulo")
    novo_run(p, "4 REFERÊNCIAS", rpr_de(modelo["ref_titulo"]), negrito=True)
    rpr_ref = rpr_de(modelo["ref"])
    for ref in ler_referencias():
        v = Paragraph(clonar("ref_vazio"), doc._body)
        v.paragraph_format.line_spacing = 1.0
        elm = clonar("ref")
        # NBR 6023: alinhadas à esquerda, entrelinha simples, sem recuo.
        par = Paragraph(elm, doc._body)
        par.alignment = WD_ALIGN_PARAGRAPH.LEFT
        par.paragraph_format.line_spacing = 1.0
        par.paragraph_format.first_line_indent = Cm(0)
        escrever(elm, ref, rpr_ref)

    for elm in originais:
        corpo_xml.remove(elm)

    doc.save(SAIDA)
    print("gerado: %s" % SAIDA)
    print("  parágrafos: %d | figuras: %d | referências: %d"
          % (len(novos), figuras, len(REFERENCIAS)))


if __name__ == "__main__":
    gerar()
